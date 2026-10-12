"""Import kartoteki dostawców z eksportu SAP LFA1 z podglądem (importers/lfa1.py, spec
2026-09-25-kartoteka-dostawcy §3): Nowi · Zmienieni · Zniknęli z SAP · Błędy, jedna kartoteka."""
import io

from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.importers.lfa1 import apply_plan, build_plan, parse_lfa1, preview
from app.models import AuditLog, Company, SapImport, Supplier

HEADER = ["Dostawca", "Klucz kraju/regionu", "Nazwa 1", "Nazwa 2", "Nazwa 3", "Nazwa 4",
          "Miasto", "Kod pocztowy", "Szukany ciąg zn.", "Ulica", "Nr ident. VAT",
          "Centralna blokada księgowania", "Utworzone przez"]

DEMOLOG = ["90000001", "DE", "DEMOLOG GMBH", "", "", "", "Hamburg", "20000", "DEMOLOG G",
            "Beispielweg 1", "DE123", "", "DEMOUSER"]
MEDDEMO = ["90000002", "FR", "MEDDEMO CARDIO", "EUROPE GMBH", "", "", "Cologne", "00000",
          "MEDDEMO COL", "", "", "X", "DEMOUSER"]
ROWS = [
    DEMOLOG,
    MEDDEMO,
    ["", "CN", "BEZ KODU", "", "", "", "", "", "", "", "", "", ""],          # błąd: brak kodu
    ["10000500", "CN", "", "", "", "", "", "", "", "", "", "", ""],          # błąd: brak nazwy
    DEMOLOG,                                                                # błąd: powtórzony
    [None] * 13,                                                             # pusty — pomijany
]


def _xlsx(rows=ROWS, header=HEADER) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _plan(db, content=None):
    return build_plan(db, *parse_lfa1(content or _xlsx()))


def test_parse_rows_and_row_errors():
    rows, errors = parse_lfa1(_xlsx())
    assert [r["sap_code"] for r in rows] == ["90000001", "90000002", "90000001"]
    assert rows[1]["name"] == "MEDDEMO CARDIO EUROPE GMBH"   # Nazwa 1 + Nazwa 2
    assert rows[0]["address"] == "Beispielweg 1, 20000 Hamburg, DE"
    assert (rows[0]["vat"], rows[0]["sap_status"], rows[1]["sap_status"]) == ("DE123", "active", "blocked")
    assert [(e["row"], e["reason"]) for e in errors] == [(4, "brak kodu SAP"), (5, "brak nazwy")]


def test_first_import_new_and_errors_then_idempotent(db_session):
    plan = _plan(db_session)
    view = preview(plan)
    assert view["counts"] == {"total": 5, "new": 2, "changed": 0, "disappeared": 0, "errors": 3}
    assert "powtórzony" in view["errors"][-1]["reason"] and view["errors"][-1]["row"] == 6
    assert db_session.scalar(select(Supplier).where(Supplier.sap_code == "90000001")) is None  # podgląd

    log = apply_plan(db_session, plan, full_export=False, user=None, filename="lfa1.xlsx")
    db_session.commit()
    stored = db_session.scalars(select(Supplier).where(Supplier.sap_code == "90000002")).one()
    assert (stored.client_company_id, stored.sap_status, stored.city) == (None, "blocked", "Cologne")
    assert log.kind == "lfa1" and log.counts["new"] == 2 and len(log.errors) == 3

    again = preview(_plan(db_session))["counts"]
    assert (again["new"], again["changed"], again["disappeared"]) == (0, 0, 0)


def test_changed_field_by_field_with_audit(db_session):
    db_session.add(Supplier(name="DEMOLOG OLD", sap_code="90000001", country="PL",
                            street="Beispielweg 1", vat="DE123", note="ręczna notatka"))
    db_session.commit()
    plan = _plan(db_session)
    change = preview(plan)["changed"][0]
    fields = {c["field"]: (c["old"], c["new"]) for c in change["changes"]}
    assert fields["name"] == ("DEMOLOG OLD", "DEMOLOG GMBH")
    assert fields["country"] == ("PL", "DE")
    assert "street" not in fields and "vat" not in fields   # bez zmian — nie pokazujemy

    apply_plan(db_session, plan, full_export=False, user=None, filename="lfa1.xlsx")
    db_session.commit()
    supplier = db_session.scalars(select(Supplier).where(Supplier.sap_code == "90000001")).one()
    assert (supplier.name, supplier.note) == ("DEMOLOG GMBH", "ręczna notatka")  # dane aplikacji nietknięte
    audit = db_session.scalars(select(AuditLog).where(
        AuditLog.entity_type == "suppliers", AuditLog.entity_id == supplier.id,
        AuditLog.field == "name")).one()
    assert (audit.old_value, audit.new_value, audit.note) == ("DEMOLOG OLD", "DEMOLOG GMBH",
                                                              "import SAP lfa1.xlsx")


def test_disappeared_deactivated_only_with_full_export(db_session):
    db_session.add(Supplier(name="STARY", sap_code="10000001"))
    db_session.commit()
    plan = _plan(db_session)
    assert [d["sap_code"] for d in preview(plan)["disappeared"]] == ["10000001"]
    apply_plan(db_session, plan, full_export=False, user=None, filename="czesc.xlsx")
    db_session.commit()
    old = db_session.scalars(select(Supplier).where(Supplier.sap_code == "10000001")).one()
    assert old.sap_status == "active"          # plik częściowy nie dezaktywuje

    log = apply_plan(db_session, _plan(db_session), full_export=True, user=None, filename="pelny.xlsx")
    db_session.commit()
    db_session.refresh(old)
    assert old.sap_status == "inactive_in_sap" and log.counts["deactivated"] == 1
    assert preview(_plan(db_session))["counts"]["disappeared"] == 0   # już oznaczony

    # wraca w pliku bez kolumn blokad → znów aktywny
    back = _xlsx([["10000001", "PL", "STARY", "", "", "", "", "", "", "", ""]], HEADER[:11])
    change = preview(_plan(db_session, back))["changed"][0]["changes"]
    assert {"field": "sap_status", "old": "inactive_in_sap", "new": "active"} in change


def test_fills_sap_code_on_catalog_record_not_on_client_sender(db_session):
    borealis = db_session.scalar(select(Company).where(Company.code == "BOREALIS"))
    db_session.add_all([Supplier(name="DEMOLOG GMBH"),
                        Supplier(name="DEMOLOG GMBH", client_company_id=borealis.id)])
    db_session.commit()
    plan = _plan(db_session)
    assert preview(plan)["counts"]["new"] == 1          # tylko MEDDEMO; DEMOLOG dopasowany po nazwie
    apply_plan(db_session, plan, full_export=True, user=None, filename="lfa1.xlsx")
    db_session.commit()
    catalog = db_session.scalars(select(Supplier).where(
        Supplier.name == "DEMOLOG GMBH", Supplier.client_company_id.is_(None))).one()
    assert catalog.sap_code == "90000001"
    sender = db_session.scalars(select(Supplier).where(Supplier.client_company_id == borealis.id)).one()
    assert (sender.sap_code, sender.sap_status) == ("", "active")   # nadawca klienta nietknięty


def test_endpoint_preview_commit_and_errors_xlsx(client, admin_headers, db_session):
    files = {"file": ("lfa1.xlsx", _xlsx(), "application/vnd.ms-excel")}
    dry = client.post("/api/import/suppliers-lfa1", headers=admin_headers, files=files)
    assert dry.status_code == 200 and dry.json()["dry_run"] is True
    assert dry.json()["counts"]["new"] == 2 and len(dry.json()["errors"]) == 3

    done = client.post("/api/import/suppliers-lfa1?dry_run=false&full_export=true",
                       headers=admin_headers, files=files).json()
    assert done["dry_run"] is False and done["deactivated"] == 0
    assert db_session.get(SapImport, done["import_id"]).counts["full_export"] is True

    report = client.get(f"/api/import/sap-imports/{done['import_id']}/errors.xlsx",
                        headers=admin_headers)
    assert report.status_code == 200
    rows = list(load_workbook(io.BytesIO(report.content)).active.iter_rows(values_only=True))
    assert rows[0] == ("Wiersz", "Kod SAP", "Nazwa", "Powód") and len(rows) == 4
    assert client.get("/api/import/sap-imports/999999/errors.xlsx",
                      headers=admin_headers).status_code == 404
