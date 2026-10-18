"""Instrukcje Excel / Power BI / Power Automate (docs/MICROSOFT-365.md) opisują dokładnie kolumny eksportu
kolejki — zmiana nagłówka w eksporcie bez poprawki instrukcji (i zapytania Power Query) zatrzyma CI."""
import io
import pathlib

from openpyxl import load_workbook

from .conftest import login

DOC = pathlib.Path(__file__).resolve().parents[2] / "docs" / "MICROSOFT-365.md"


def test_doc_lists_every_export_column_and_query_uses_real_names(client):
    r = client.get("/api/containers/export/xlsx", headers=login(client))
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.title == "KOLEJKA"
    headers = [c.value for c in ws[1]]
    text = DOC.read_text(encoding="utf-8")
    table = [line.split("|")[1].strip() for line in text.splitlines()
             if line.startswith("| ") and line.count("|") >= 4]
    assert [h for h in table if h in headers] == headers           # wszystkie, w tej kolejności
    query = text[text.index("```powerquery"):text.index("```", text.index("```powerquery") + 3)]
    for used in ("DATA AWIZACJI", "ETA", "ATD", "DNI WOLNE OD DEMURRAGE", "DOKUMENTY", "OPÓŹNIONY"):
        assert f'"{used}"' in query and used in headers
